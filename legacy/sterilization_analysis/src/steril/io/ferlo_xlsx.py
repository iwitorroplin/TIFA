"""Etapa 1 - Lector de los .xlsx de Ferlo.

Los ficheros reales tienen cuatro columnas:

    Fecha dd/MM/yyyy | Hora H:mm:ss | Valor °C | Valor bar

Exportaciones del programa antiguo anaden una quinta, `Tiempo`, que la etapa 2
usa solo como comprobacion cruzada. El lector la conserva si esta y no la echa
de menos si no esta.

`read_only=True` es obligatorio: sin el, openpyxl construye en memoria todas las
celdas del fichero y tarda un orden de magnitud mas.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .base import RawTable, register_reader, sha256_of


class FerloXlsxReader:
    name = "ferlo_xlsx"
    extensions = (".xlsx", ".xlsm")

    def can_read(self, path: Path) -> bool:
        return path.suffix.lower() in self.extensions

    def read(self, path: Path, *, sheet_index: int = 0, **_: Any) -> RawTable:
        path = Path(path)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            try:
                ws = wb.worksheets[sheet_index]
            except IndexError:
                raise ValueError(
                    f"{path.name}: no existe la hoja {sheet_index} "
                    f"(el fichero tiene {len(wb.worksheets)})"
                ) from None

            filas = ws.iter_rows(values_only=True)
            try:
                cabecera = next(filas)
            except StopIteration:
                return RawTable(path, ws.title, [], [], sha256_of(path))

            headers = [("" if c is None else str(c)).strip() for c in cabecera]
            ancho = len(headers)
            rows: list[tuple[Any, ...]] = []
            for fila in filas:
                # Excel devuelve filas vacias al final del area usada.
                if all(v is None for v in fila):
                    continue
                if len(fila) < ancho:
                    fila = fila + (None,) * (ancho - len(fila))
                rows.append(tuple(fila[:ancho]))
        finally:
            wb.close()

        return RawTable(
            source_path=path,
            sheet_name=ws.title,
            headers=headers,
            rows=rows,
            source_sha256=sha256_of(path),
        )


register_reader(FerloXlsxReader())
